from __future__ import annotations

import tomllib
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from .models import (
    Member,
    ProjectMatch,
    ProjectRoute,
    ProjectRule,
    ServiceRule,
    Ticket,
)
from .text import normalize, unaccent


class RuleError(ValueError):
    pass


_PROJECT_FIELDS = {
    "subscriber_name", "customer_name", "labor_address", "connection_address",
    "labor_province", "notes",
}

# Tên rút gọn đã có trong các dòng lịch sử trên Google Sheet. Giá trị được
# chuẩn hóa về tên đầy đủ đang khai báo tại tiêu đề file Excel trước khi xác
# định nhóm hoặc giao lại phiếu.
_MEMBER_NAME_ALIASES = {
    normalize("Lý Bích Hằng"): "Lý Thị Bích Hằng",
}


def load_project_rules(path: str | Path | None) -> list[ProjectRule]:
    """Load named high-priority routing rules from a small editable TOML file."""
    if path is None:
        return []
    project_path = Path(path)
    if not project_path.exists():
        raise RuleError(f"Không tìm thấy file quy tắc dự án: {project_path}")
    with project_path.open("rb") as handle:
        data = tomllib.load(handle)
    rules: list[ProjectRule] = []
    for index, item in enumerate(data.get("projects", []), start=1):
        name = str(item.get("name", "")).strip()
        raw_contains = item.get("contains", "")
        contains_values = (
            tuple(str(value).strip() for value in raw_contains)
            if isinstance(raw_contains, list)
            else (str(raw_contains).strip(),)
        )
        contains_values = tuple(value for value in contains_values if value)
        fields = tuple(str(value).strip() for value in item.get("match_fields", []))
        if not name or not contains_values or not fields:
            raise RuleError(
                f"Quy tắc dự án số {index} phải có name, contains và match_fields"
            )
        invalid_fields = set(fields) - _PROJECT_FIELDS
        raw_required = item.get("required_contains", {})
        if not isinstance(raw_required, dict):
            raise RuleError(f"Quy tắc {name} phải khai báo required_contains dạng bảng")
        required_contains = tuple(
            (str(field).strip(), str(value).strip())
            for field, value in raw_required.items()
            if str(field).strip() and str(value).strip()
        )
        invalid_fields.update(field for field, _ in required_contains if field not in _PROJECT_FIELDS)
        route_field = str(item.get("route_field", "")).strip()
        if route_field and route_field not in _PROJECT_FIELDS:
            invalid_fields.add(route_field)
        if invalid_fields:
            raise RuleError(
                f"Quy tắc {name} dùng trường không hợp lệ: {', '.join(sorted(invalid_fields))}"
            )
        routes = tuple(
            ProjectRoute(
                assignee=str(route.get("assignee", "")).strip(),
                locations=tuple(
                    str(value).strip() for value in route.get("locations", []) if str(value).strip()
                ),
                subscriber_prefixes=tuple(
                    str(value).strip() for value in route.get("subscriber_prefixes", [])
                    if str(value).strip()
                ),
            )
            for route in item.get("routes", [])
        )
        fixed_assignee = str(item.get("fixed_assignee", "")).strip()
        if not fixed_assignee and (not route_field or not routes):
            raise RuleError(
                f"Quy tắc {name} phải có fixed_assignee hoặc route_field và routes"
            )
        if any(not route.assignee or not route.locations for route in routes):
            raise RuleError(f"Quy tắc {name} có tuyến địa bàn chưa đầy đủ")
        # Expand multiple identifying phrases into equivalent rules that share
        # the same project name and routing table. This keeps one authoritative
        # set of province routes for aliases such as the BTC organisations.
        rules.extend(
            ProjectRule(
                name=name,
                contains=contains,
                match_fields=fields,
                priority=int(item.get("priority", 0)),
                fixed_assignee=fixed_assignee,
                route_field=route_field,
                routes=routes,
                required_contains=required_contains,
            )
            for contains in contains_values
        )
    return sorted(rules, key=lambda rule: rule.priority, reverse=True)


def match_project_rule(
    project_rules: list[ProjectRule], ticket: Ticket
) -> ProjectMatch | None:
    """Return a named project override, or fail safely for an unknown project area."""
    for rule in project_rules:
        values = [unaccent(getattr(ticket, field, "")) for field in rule.match_fields]
        if not any(unaccent(rule.contains) in value for value in values):
            continue
        if any(
            unaccent(expected) not in unaccent(getattr(ticket, field, ""))
            for field, expected in rule.required_contains
        ):
            continue
        if rule.fixed_assignee:
            return ProjectMatch(rule.name, rule.fixed_assignee)
        route_fields = [rule.route_field]
        # Every routed project uses Địa chỉ LĐ first. If that text does not
        # identify a configured province/city, fall back to the explicit
        # Tỉnh LĐ column from the OneBSS ticket grid.
        if rule.route_field == "labor_address":
            route_fields.append("labor_province")
        for route_field in route_fields:
            route_value = unaccent(getattr(ticket, route_field, ""))
            for route in rule.routes:
                location_matches = any(
                    unaccent(location) in route_value for location in route.locations
                )
                prefix_matches = (
                    not route.subscriber_prefixes
                    or any(
                        normalize(ticket.subscriber_id).startswith(normalize(prefix))
                        for prefix in route.subscriber_prefixes
                    )
                )
                if location_matches and prefix_matches:
                    return ProjectMatch(rule.name, route.assignee)
        shown_value = getattr(ticket, rule.route_field, "")
        raise RuleError(
            f"{rule.name}: không xác định được người nhận từ "
            f"{rule.route_field}={shown_value!r}, "
            f"labor_province={ticket.labor_province!r}"
        )
    return None


def canonical_member_name(rules: list[ServiceRule], name: str) -> str:
    """Resolve a Sheet name to the canonical employee spelling from Excel.

    Manual Sheet entries occasionally omit or mistype Vietnamese tone marks
    (for example ``Đăng Văn Minh`` instead of ``Đặng Văn Minh``). Exact
    Unicode-normalized matches remain preferred; an accent-insensitive match
    is accepted only when it identifies exactly one configured employee.
    """
    names = {
        member.name
        for rule in rules
        for member in rule.members
        if member.name.strip()
    }
    wanted = normalize(_MEMBER_NAME_ALIASES.get(normalize(name), name))
    exact = {candidate for candidate in names if normalize(candidate) == wanted}
    if len(exact) == 1:
        return exact.pop()
    accentless = {
        candidate for candidate in names if unaccent(candidate) == unaccent(name)
    }
    if len(accentless) == 1:
        return accentless.pop()
    if len(accentless) > 1:
        raise RuleError(
            f"Tên nhân viên {name} không rõ ràng trong file quy tắc: "
            + ", ".join(sorted(accentless))
        )
    raise RuleError(f"Nhân viên {name} không có trong nhóm của file quy tắc")


def member_group(
    rules: list[ServiceRule], name: str, service_rule: ServiceRule | None = None
) -> str:
    canonical = canonical_member_name(rules, name)
    wanted = normalize(canonical)
    if service_rule is not None:
        service_groups = {
            member.group
            for member in service_rule.members
            if normalize(member.name) == wanted
        }
        if len(service_groups) == 1:
            return service_groups.pop()
        if len(service_groups) > 1:
            raise RuleError(
                f"Nhân viên {name} thuộc nhiều nhóm trong cùng quy tắc dịch vụ"
            )
    groups = {
        member.group
        for rule in rules
        for member in rule.members
        if normalize(member.name) == wanted
    }
    if len(groups) > 1:
        raise RuleError(f"Nhân viên {name} xuất hiện ở nhiều nhóm: {', '.join(sorted(groups))}")
    return groups.pop()


def member_score_key(group: str, name: str) -> str:
    """Stable score identity for employees who belong to more than one group."""
    return f"{normalize(unaccent(group))}::{normalize(name)}"


def _load_group_rules(workbook) -> list[ServiceRule]:
    """Expand percentage-by-group rules through the OneBSS service map."""
    rule_sheet = workbook[workbook.sheetnames[0]]
    headers = [
        str(rule_sheet.cell(2, col).value or "").strip()
        for col in range(1, rule_sheet.max_column + 1)
    ]
    header_columns = {
        unaccent(header): index + 1
        for index, header in enumerate(headers) if header
    }
    group_column = header_columns["group dv"]
    points_column = header_columns.get("diem quy doi")
    condition_column = header_columns.get("quy tac")
    api_column = header_columns.get("gui api")
    if not points_column or not condition_column:
        raise RuleError(
            "File Giao phiếu dạng Group DV phải có cột Điểm quy đổi và Quy tắc ở dòng 2"
        )
    member_start = max(group_column, points_column, condition_column) + 1
    member_end = api_column - 1 if api_column else rule_sheet.max_column
    if member_start > member_end:
        raise RuleError("File Giao phiếu dạng Group DV không có cột nhân sự")

    group_by_column: dict[int, str] = {}
    current_group = ""
    for col in range(member_start, member_end + 1):
        label = str(rule_sheet.cell(1, col).value or "").strip()
        if label:
            current_group = label
        group_by_column[col] = current_group or "Nhóm 1"

    grouped_rules: dict[tuple[str, Decimal], tuple[int, str, Decimal, tuple[Member, ...], bool, bool, bool]] = {}
    for row in range(3, rule_sheet.max_row + 1):
        group_name = str(rule_sheet.cell(row, group_column).value or "").strip()
        if not group_name:
            continue
        raw_points = rule_sheet.cell(row, points_column).value
        try:
            points = Decimal(str(raw_points or 0))
        except Exception as error:
            raise RuleError(f"Điểm quy đổi không hợp lệ tại dòng {row}: {raw_points!r}") from error
        condition = str(rule_sheet.cell(row, condition_column).value or "").strip()
        members: list[Member] = []
        share_total = Decimal(0)
        for col in range(member_start, member_end + 1):
            raw_share = rule_sheet.cell(row, col).value
            if raw_share in (None, ""):
                continue
            try:
                text = str(raw_share).strip()
                is_percent = text.endswith("%")
                share = Decimal(text.removesuffix("%"))
                if is_percent or (isinstance(raw_share, (int, float)) and share > 1):
                    share /= 100
            except Exception as error:
                raise RuleError(
                    f"Tỷ lệ nhân sự không hợp lệ tại dòng {row}, cột {col}: {raw_share!r}"
                ) from error
            if share < 0 or share > 1:
                raise RuleError(
                    f"Tỷ lệ nhân sự phải từ 0% đến 100% tại dòng {row}, cột {col}"
                )
            if share:
                name = headers[col - 1]
                if not name:
                    raise RuleError(f"Thiếu tên nhân sự ở cột {col}, dòng {row}")
                members.append(Member(name, group_by_column[col], "Chính", share))
                share_total += share
        if not members:
            raise RuleError(f"Nhóm dịch vụ {group_name} không có nhân sự được phân tỷ lệ")
        if abs(share_total - Decimal(1)) > Decimal("0.000001"):
            raise RuleError(
                f"Tổng tỷ lệ của nhóm dịch vụ {group_name} tại dòng {row} là "
                f"{share_total * 100}%, phải bằng 100%"
            )
        api_value = (
            unaccent(str(rule_sheet.cell(row, api_column).value or "").strip())
            if api_column else ""
        )
        send_to_api = api_value in {"co", "yes", "true", "1", "x", "gui", "gui api"}
        policy = unaccent(condition)
        group_key = (unaccent(group_name), points)
        if group_key in grouped_rules:
            raise RuleError(
                f"Nhóm DV {group_name} có nhiều quy tắc cùng {points} điểm; "
                "không thể ánh xạ dịch vụ một cách duy nhất"
            )
        grouped_rules[group_key] = (
            row, group_name, points, tuple(members),
            "khong tinh diem" not in policy,
            "khong dua vao danh sach" not in policy,
            send_to_api,
        )

    if "group dv" not in {
        unaccent(str(rule_sheet.cell(2, col).value or "").strip())
        for col in range(1, rule_sheet.max_column + 1)
    }:
        raise RuleError("Thiếu cột Group DV")
    if len(workbook.worksheets) < 2:
        raise RuleError("File Giao phiếu dạng Group DV cần có sheet Group DV để ánh xạ dịch vụ")
    map_sheet = next(
        (sheet for sheet in workbook.worksheets[1:]
         if unaccent(sheet.title) == "group dv"),
        None,
    )
    if map_sheet is None:
        raise RuleError("Không tìm thấy sheet Group DV để ánh xạ dịch vụ")
    map_headers = {
        unaccent(str(map_sheet.cell(1, col).value or "").strip()): col
        for col in range(1, map_sheet.max_column + 1)
    }
    required = ("dich vu", "dich vu tren google sheet", "diem quy doi", "group dv")
    if any(header not in map_headers for header in required):
        raise RuleError(
            "Sheet Group DV cần có các cột Dịch vụ, Dịch vụ trên Google Sheet, "
            "Điểm quy đổi và Group DV"
        )
    rules: list[ServiceRule] = []
    signatures: set[tuple] = set()
    for row in range(2, map_sheet.max_row + 1):
        service = str(map_sheet.cell(row, map_headers["dich vu"]).value or "").strip()
        group_name = str(map_sheet.cell(row, map_headers["group dv"]).value or "").strip()
        if not service and not group_name:
            continue
        if not service or not group_name:
            raise RuleError(f"Thiếu Dịch vụ hoặc Group DV tại sheet Group DV, dòng {row}")
        raw_points = map_sheet.cell(row, map_headers["diem quy doi"]).value
        try:
            points = Decimal(str(raw_points or 0))
        except Exception as error:
            raise RuleError(f"Điểm quy đổi không hợp lệ ở sheet Group DV, dòng {row}") from error
        grouped = grouped_rules.get((unaccent(group_name), points))
        if grouped is None:
            raise RuleError(
                f"Không tìm thấy quy tắc nhóm {group_name} với {points} điểm "
                f"(sheet Group DV, dòng {row})"
            )
        base_row, _base_group_name, _base_points, members, count_points, write_to_sheet, send_to_api = grouped
        base_condition = str(rule_sheet.cell(base_row, condition_column).value or "").strip()
        map_condition_column = map_headers.get("quy tac group")
        map_condition = (
            str(map_sheet.cell(row, map_condition_column).value or "").strip()
            if map_condition_column else ""
        )
        if base_condition and map_condition and unaccent(base_condition) != unaccent(map_condition):
            raise RuleError(
                f"Quy tắc group không khớp tại sheet Group DV, dòng {row}"
            )
        condition = base_condition or map_condition
        raw_sheet_service = map_sheet.cell(
            row, map_headers["dich vu tren google sheet"]
        ).value
        sheet_service = str(raw_sheet_service or "").strip() or service
        signature = (
            normalize(service), normalize(sheet_service), points, normalize(condition),
            tuple((normalize(m.name), normalize(m.group), m.target_share) for m in members),
        )
        if signature in signatures:
            raise RuleError(f"Ánh xạ dịch vụ bị lặp tại sheet Group DV, dòng {row}")
        signatures.add(signature)
        rules.append(ServiceRule(
            base_row, service, points, condition, members, sheet_service,
            count_points, write_to_sheet, send_to_api,
        ))
    if not rules:
        raise RuleError("Sheet Group DV không có dòng ánh xạ dịch vụ hợp lệ")
    return rules


def load_rules(path: str | Path) -> list[ServiceRule]:
    workbook = load_workbook(Path(path), data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    headers = [str(sheet.cell(2, col).value or "").strip() for col in range(1, sheet.max_column + 1)]
    header_columns = {unaccent(header): index + 1 for index, header in enumerate(headers) if header}
    if "group dv" in header_columns:
        return _load_group_rules(workbook)
    service_column = header_columns.get("dich vu")
    sheet_service_column = header_columns.get("dich vu tren google sheet")
    points_column = header_columns.get("diem quy doi")
    condition_column = header_columns.get("quy tac")
    api_column = header_columns.get("gui api")
    if not service_column or not points_column or not condition_column:
        raise RuleError(
            "File quy tắc phải có các cột Dịch vụ, Điểm quy đổi và Quy tắc ở dòng 2"
        )
    member_start_column = max(
        service_column, sheet_service_column or 0, points_column, condition_column
    ) + 1
    group_by_column: dict[int, str] = {}
    current_group = ""
    for col in range(member_start_column, sheet.max_column + 1):
        if col == api_column:
            continue
        label = str(sheet.cell(1, col).value or "").strip()
        if label:
            current_group = label
        group_by_column[col] = current_group or "Nhóm 1"

    rules: list[ServiceRule] = []
    signatures: set[tuple] = set()
    for row in range(3, sheet.max_row + 1):
        service = str(sheet.cell(row, service_column).value or "").strip()
        if not service:
            continue
        raw_sheet_service = (
            sheet.cell(row, sheet_service_column).value if sheet_service_column else service
        )
        sheet_service = str(raw_sheet_service or "").strip() or service
        raw_points = sheet.cell(row, points_column).value
        points = Decimal(str(raw_points or 0))
        condition = str(sheet.cell(row, condition_column).value or "").strip()
        policy = unaccent(condition)
        count_points = "khong tinh diem" not in policy
        write_to_sheet = "khong dua vao danh sach" not in policy
        api_value = (
            unaccent(str(sheet.cell(row, api_column).value or "").strip())
            if api_column else ""
        )
        send_to_api = (
            True if api_column is None
            else api_value in {"co", "yes", "true", "1", "x", "gui", "gui api"}
        )
        members: list[Member] = []
        for col in range(member_start_column, sheet.max_column + 1):
            if col == api_column:
                continue
            role = str(sheet.cell(row, col).value or "").strip()
            if role:
                members.append(Member(headers[col - 1], group_by_column[col], role))
        signature = (
            normalize(service), normalize(sheet_service), points, normalize(condition),
            tuple((normalize(m.name), normalize(m.group), normalize(m.role)) for m in members),
            count_points, write_to_sheet, send_to_api,
        )
        if signature in signatures:
            continue
        signatures.add(signature)
        rules.append(ServiceRule(
            row, service, points, condition, tuple(members), sheet_service,
            count_points, write_to_sheet, send_to_api,
        ))
    return rules


def load_score_member_names(path: str | Path) -> tuple[str, ...]:
    """Return Excel member columns that participate in monthly point totals."""
    rules = load_rules(path)
    scored_names = {
        normalize(member.name)
        for rule in rules
        if rule.count_points
        for member in rule.members
    }
    workbook = load_workbook(Path(path), data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    headers = [
        str(sheet.cell(2, col).value or "").strip()
        for col in range(1, sheet.max_column + 1)
    ]
    workbook.close()
    return tuple(name for name in headers if normalize(name) in scored_names)


def _condition_matches(rule: ServiceRule, ticket: Ticket) -> bool:
    condition = unaccent(rule.condition)
    if not condition:
        return True
    if "cot vip xu ly" in condition:
        vip = unaccent(ticket.vip_status)
        if "giam sat" in condition:
            return "giam sat" in vip
        if "xu ly" in condition:
            return "xu ly" in vip
        return False
    if "loai kenh" in condition:
        channel = unaccent(ticket.channel_type)
        if "noi tinh" in condition and "lien tinh" not in condition:
            return "noi tinh" in channel
        return not channel or "lien tinh" in channel
    if "loai hd" in condition:
        contract = unaccent(ticket.contract_type)
        candidates = condition.split(":", 1)[-1]
        if any(unaccent(item).strip() in contract for item in candidates.split(",")):
            return True
        # Different deployments can phrase the same Loại HĐ differently.
        # This comparison uses only Loại HĐ from OneBSS; it never uses the
        # transaction code as a substitute.
        action_names = ("gia han", "cham dut", "lap dat moi", "chuyen quyen")
        if any(action in condition and action in contract for action in action_names):
            return True
        return False
    if "kieu lap dat" in condition:
        installation = unaccent(ticket.installation_type)
        candidates = condition.split(":", 1)[-1]
        return any(
            unaccent(item).strip() in installation
            for item in candidates.split(",")
        )
    return True


def match_rule(rules: list[ServiceRule], ticket: Ticket) -> ServiceRule:
    service = normalize(ticket.service_type or ticket.service)
    candidates = [rule for rule in rules if normalize(rule.service) == service]
    matches = [rule for rule in candidates if _condition_matches(rule, ticket)]
    if not matches:
        raise RuleError(
            f"Không tìm thấy quy tắc cho dịch vụ={ticket.service_type or ticket.service!r}, "
            f"loại HĐ={ticket.contract_type!r}, kiểu lắp đặt={ticket.installation_type!r}, "
            f"loại kênh={ticket.channel_type!r}"
        )
    if len(matches) > 1:
        # Prefer a specific condition over an unconditional row.
        matches.sort(key=lambda rule: bool(rule.condition), reverse=True)
        first = matches[0]
        same = [m for m in matches if bool(m.condition) == bool(first.condition)]
        if len(same) > 1 and any(m.points != first.points for m in same):
            raise RuleError(f"Có nhiều quy tắc không phân biệt được cho {ticket.service_type or ticket.service!r}")
    return matches[0]


def voice_brandname_mode(rule: ServiceRule, ticket: Ticket) -> str:
    """Return the mandatory Voice Brandname route from the current VIP state."""
    if normalize(rule.service) != normalize("Voice Brandname"):
        return ""
    vip = unaccent(ticket.vip_status)
    if "giam sat" in vip:
        return "giam_sat"
    if "xu ly" in vip:
        return "xu_ly"
    raise RuleError(
        "Voice Brandname không xác định được VIP xử lý "
        f"(giá trị hiện tại: {ticket.vip_status!r})"
    )


def eligible_members(
    rule: ServiceRule,
    ticket: Ticket,
    use_backups: bool,
    excluded_members: tuple[str, ...] = (),
) -> dict[str, list[str]]:
    excluded = {normalize(name) for name in excluded_members}
    voice_mode = voice_brandname_mode(rule, ticket)
    if voice_mode == "giam_sat":
        selected_names = ["Lê Đức Tuấn"]
        selected_names = [
            name for name in selected_names if normalize(name) not in excluded
        ]
        if not selected_names:
            raise RuleError("Lê Đức Tuấn đang được loại trừ do nghỉ phép")
        return {"Nhóm 2": selected_names}
    if voice_mode == "xu_ly":
        selected_names = [
            "Đỗ Thị Thu Trang", "Ngô Thùy Trang", "Nguyễn Thị Thu Trang",
        ]
        selected_names = [
            name for name in selected_names if normalize(name) not in excluded
        ]
        if not selected_names:
            raise RuleError(
                "Toàn bộ nhân sự xử lý Voice Brandname đang nghỉ phép"
            )
        return {"Nhóm 2": selected_names}

    selected = [
        member for member in rule.members
        if normalize(member.name) not in excluded
    ]
    if not use_backups and any(unaccent(member.role) == "chinh" for member in selected):
        selected = [member for member in selected if unaccent(member.role) == "chinh"]
    groups: dict[str, list[str]] = {}
    for member in selected:
        groups.setdefault(member.group, []).append(member.name)
    if not groups:
        suffix = " sau khi loại nhân sự nghỉ phép" if excluded else ""
        raise RuleError(
            f"Quy tắc dòng {rule.row_number} không có nhân viên phù hợp{suffix}"
        )
    return groups
