from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

import pytest

from ats_onebss.models import ProjectRoute, ProjectRule, Ticket
from ats_onebss.models import Member, ServiceRule
from ats_onebss.rules import (
    RuleError,
    canonical_member_name,
    eligible_members,
    load_project_rules,
    load_rules,
    load_score_member_names,
    match_project_rule,
    match_rule,
    member_group,
)
from ats_onebss.config import load_config
from ats_onebss.text import normalize, unaccent


def make_rules(tmp_path):
    path = tmp_path / "rules.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append([None, None, None, None, "Nhóm 1", None, "Nhóm 2"])
    sheet.append(["STT", "Dịch vụ", "Điểm quy đổi", "Quy tắc", "An", "Bình", "Hoa"])
    sheet.append([1, "Tên miền Việt Nam", 14, "Loại HĐ: Gia hạn, Chấm dứt", "Chính", "Phụ", None])
    sheet.append([2, "Tên miền Việt Nam", 29, "Loại HĐ: Lắp đặt mới, Chuyển quyền", "Chính", None, "Chính"])
    book.save(path)
    return path


def test_loads_separate_google_sheet_service_and_shifted_members(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([None, None, None, None, None, "Nhóm 1"])
    sheet.append([
        "STT", "Dịch vụ", "Dịch vụ trên Google Sheet", "Điểm quy đổi",
        "Quy tắc", "An",
    ])
    sheet.append([1, "ISDN 30B+D cáp đồng", "B-FONE", 17, None, "Chính"])
    path = tmp_path / "rules-new.xlsx"
    workbook.save(path)

    rule = load_rules(path)[0]
    assert rule.service == "ISDN 30B+D cáp đồng"
    assert rule.sheet_service == "B-FONE"
    assert rule.members[0].name == "An"


def test_api_column_is_opt_in_and_not_treated_as_a_member(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([None, None, None, None, "Nhóm 1", None])
    sheet.append(["STT", "Dịch vụ", "Điểm quy đổi", "Quy tắc", "An", "Gửi API"])
    sheet.append([1, "Fiber", 17, "", "Chính", "Có"])
    sheet.append([2, "MetroNet LT", 27, "", "Chính", None])
    path = tmp_path / "rules-api.xlsx"
    workbook.save(path)

    rules = load_rules(path)

    assert rules[0].send_to_api is True
    assert [member.name for member in rules[0].members] == ["An"]
    assert rules[1].send_to_api is False
    assert [member.name for member in rules[1].members] == ["An"]


def test_workbook_without_api_column_preserves_legacy_api_behavior(tmp_path):
    rules = load_rules(make_rules(tmp_path))

    assert all(rule.send_to_api for rule in rules)


def test_current_workbook_contains_leasedline_ge():
    workbook = Path(__file__).parents[1] / "Giao phiếu.xlsx"
    rule = next(item for item in load_rules(workbook) if item.service == "Leasedline GE")
    assert rule.sheet_service == "Kênh Thuê Riêng"
    assert rule.points == Decimal("42")
    assert {member.name for member in rule.members} == {
        "Vũ Thế Ninh", "Đào Anh Vũ", "Đoàn Hải Hà", "Nguyễn Duy Thành",
    }


def test_current_workbook_excludes_international_voice_from_points_and_sheet():
    workbook = Path(__file__).parents[1] / "Giao phiếu.xlsx"
    rule = next(item for item in load_rules(workbook) if item.service == "Thoại quốc tế")

    assert rule.points == Decimal("0")
    assert rule.count_points is False
    assert rule.write_to_sheet is False
    assert [normalize(member.name) for member in rule.members] == [
        normalize("Trần Mạnh Cường")
    ]


def test_policy_text_controls_points_and_google_sheet_output(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([None, None, None, None, "Nhóm 1"])
    sheet.append(["STT", "Dịch vụ", "Điểm quy đổi", "Quy tắc", "An"])
    sheet.append([
        1,
        "Thoại quốc tế",
        99,
        "Không tính điểm, không đưa vào danh sách (Vẫn phân phiếu)",
        "Chính",
    ])
    path = tmp_path / "rules.xlsx"
    workbook.save(path)

    rule = load_rules(path)[0]
    assert rule.points == Decimal("99")
    assert rule.count_points is False
    assert rule.write_to_sheet is False


def test_load_and_match_contract_rule(tmp_path):
    rules = load_rules(make_rules(tmp_path))
    ticket = Ticket("GD1", "TB1", "Dịch vụ CNTT", "Tên miền Việt Nam", "Lắp đặt mới")
    rule = match_rule(rules, ticket)
    assert rule.points == Decimal("29")
    assert rule.row_number == 4


def test_contract_rule_matches_onebss_renewal_wording(tmp_path):
    rules = load_rules(make_rules(tmp_path))
    ticket = Ticket(
        "GD1", "TB1", "Dịch vụ CNTT", "Tên miền Việt Nam",
        "Gia hạn và thay đổi gói cước CNTT/GTGT",
    )
    rule = match_rule(rules, ticket)
    assert rule.points == Decimal("14")


def test_current_domain_rules_match_installation_type_not_contract_type():
    rules = load_rules(Path(__file__).parents[1] / "Giao phiếu.xlsx")
    expected = {
        "Gia hạn và thay đổi gói cước CNTT/GTGT": (
            Decimal("14"), {"Lý Thị Bích Hằng", "Tạ Lê Hoa"},
        ),
        "Đặt mới dịch vụ CNTT": (
            Decimal("29"), {"Lý Thị Bích Hằng", "Tạ Lê Hoa"},
        ),
        "Chuyển quyền sử dụng CNTT": (Decimal("29"), {"Tạ Lê Hoa"}),
        "Thay đổi thông tin DV CNTT/GTGT": (Decimal("14"), {"Tạ Lê Hoa"}),
    }
    for service in ("Tên miền Việt Nam", "Tên miền Quốc tế"):
        for installation_type, (points, members) in expected.items():
            ticket = Ticket(
                "GD1", "TB1", service,
                contract_type="Giá trị Loại HĐ không dùng cho Tên miền",
                installation_type=installation_type,
            )
            rule = match_rule(rules, ticket)
            assert "kieu lap dat" in unaccent(rule.condition)
            assert rule.points == points
            assert {member.name for member in rule.members} == members


def test_domain_renewal_does_not_fall_back_to_gh_transaction_code(tmp_path):
    rules = load_rules(make_rules(tmp_path))
    ticket = Ticket(
        "CBG-GH/00013381", "domain.vn", "Dịch vụ CNTT", "Tên miền Việt Nam"
    )
    with pytest.raises(RuleError, match="Không tìm thấy quy tắc"):
        match_rule(rules, ticket)


def test_primary_members_are_preferred(tmp_path):
    rules = load_rules(make_rules(tmp_path))
    ticket = Ticket("GD1", "TB1", "Dịch vụ CNTT", "Tên miền Việt Nam", "Gia hạn")
    groups = eligible_members(match_rule(rules, ticket), ticket, use_backups=False)
    assert groups == {"Nhóm 1": ["An"]}


def test_voice_brandname_xu_ly_uses_processing_group():
    rule = ServiceRule(
        3, "Voice Brandname", Decimal("20"), "",
        (Member("Người khác", "Nhóm 1", "Chính"),),
    )
    ticket = Ticket("GD1", "TB1", "Voice Brandname", vip_status="Xu Ly")
    assert eligible_members(rule, ticket, False) == {
        "Nhóm 2": ["Đỗ Thị Thu Trang", "Ngô Thùy Trang", "Nguyễn Thị Thu Trang"]
    }


def test_voice_brandname_unknown_vip_is_rejected():
    rule = ServiceRule(
        3, "Voice Brandname", Decimal("20"), "",
        (Member("Lê Đức Tuấn", "Nhóm 2", "Chính"),),
    )
    ticket = Ticket("GD1", "TB1", "Voice Brandname", vip_status="")
    with pytest.raises(RuleError, match="không xác định được VIP xử lý"):
        eligible_members(rule, ticket, False)


def test_score_member_names_exclude_no_score_dispatcher():
    rules_file = Path(__file__).parents[1] / "Giao phiếu.xlsx"
    names = load_score_member_names(rules_file)
    assert len(names) == 17
    assert names[0] == "Lương Tuấn Thanh"
    assert not any("Cường" in name or "Cường" in name for name in names)


def test_current_workbook_has_requested_fourteen_and_four_member_groups():
    root = Path(__file__).parents[1]
    rules_file = root / "Giao phiếu.xlsx"
    rules = load_rules(rules_file)
    all_members = {member.name for rule in rules for member in rule.members}
    grouped = {
        group: {name for name in all_members if member_group(rules, name) == group}
        for group in {member_group(rules, name) for name in all_members}
    }
    group_1 = next(values for key, values in grouped.items() if "1" in key)
    group_2 = next(values for key, values in grouped.items() if "2" in key)
    assert group_1 == {
        "Lương Tuấn Thanh", "Lý Thị Bích Hằng", "Tạ Lê Hoa", "Lê Anh Tuấn",
        "Vũ Thế Ninh", "Đặng Văn Minh", "Đào Anh Vũ", "Đoàn Hải Hà",
        "Nguyễn Duy Thành", "Ngô Thị Minh Phương", "Lê Đức Vinh",
        "Trâ\u0300n Ma\u0323nh Cươ\u0300ng", "Nguyễn Hoàng Dương", "Trần Thị Thu Hằng",
    }
    assert group_2 == {
        "Lê Đức Tuấn", "Đỗ Thị Thu Trang", "Ngô Thùy Trang",
        "Nguyễn Thị Thu Trang",
    }


def test_current_config_has_requested_member_target_ratios():
    config = load_config(Path(__file__).parents[1] / "config.toml")
    assert config.member_target_ratios == {
        "Lương Tuấn Thanh": Decimal("1.10"),
        "Đào Anh Vũ": Decimal("1.10"),
        "Nguyễn Hoàng Dương": Decimal("1.10"),
        "Đoàn Hải Hà": Decimal("1.05"),
        "Nguyễn Duy Thành": Decimal("1.05"),
    }


def test_demo2_percentage_workbook_expands_service_map_and_preserves_api_flags():
    workbook = Path(__file__).parents[1] / "Giao phiếu_demo2.xlsx"
    rules = load_rules(workbook)

    assert len(rules) == 60
    sip = match_rule(rules, Ticket("GD1", "TB1", "ISDN 30B+D cáp đồng"))
    assert sip.sheet_service == "SIP"
    assert {member.name: member.target_share for member in sip.members} == {
        "Nguyễn Hoàng Dương": Decimal("0.3"),
        "Lê Đức Tuấn": Decimal("0.7"),
    }
    fiber = match_rule(rules, Ticket("GD2", "TB2", "Mega"))
    assert fiber.sheet_service == "Fiber"
    assert fiber.send_to_api is True


def test_demo2_service_map_uses_channel_and_installation_conditions():
    workbook = Path(__file__).parents[1] / "Giao phiếu_demo2.xlsx"
    rules = load_rules(workbook)

    local = match_rule(rules, Ticket("GD1", "TB1", "MegaWan", channel_type="Nội tỉnh"))
    interprovincial = match_rule(
        rules, Ticket("GD2", "TB2", "MegaWan", channel_type="Liên tỉnh")
    )
    assert local.sheet_service == "Megawan NT"
    assert local.points == Decimal("17")
    assert interprovincial.sheet_service == "Megawan LT"
    assert interprovincial.points == Decimal("27")

    domain = match_rule(
        rules,
        Ticket("GD3", "TB3", "Tên miền Việt Nam", installation_type="Thanh lý DV CNTT"),
    )
    assert domain.sheet_service == "Tên Miền Việt Nam (Triển khai)"
    assert domain.points == Decimal("14")

    voice_xu_ly = match_rule(
        rules, Ticket("GD4", "TB4", "Voice Brandname", vip_status="Xu ly")
    )
    voice_giam_sat = match_rule(
        rules, Ticket("GD5", "TB5", "Voice Brandname", vip_status="Giam sat")
    )
    assert {member.name for member in voice_xu_ly.members} == {
        "Đỗ Thị Thu Trang", "Ngô Thùy Trang", "Nguyễn Thị Thu Trang",
    }
    assert [member.name for member in voice_giam_sat.members] == ["Lê Đức Tuấn"]


def test_demo2_percentage_rows_sum_to_100_and_keep_fixed_group_roster():
    workbook = Path(__file__).parents[1] / "Giao phiếu_demo2.xlsx"
    rules = load_rules(workbook)
    names = load_score_member_names(workbook)
    group_members = {
        group: {
            member.name
            for rule in rules for member in rule.members
            if member.group == group
        }
        for group in {member.group for rule in rules for member in rule.members}
    }

    assert len(group_members["Nhóm 1"]) == 14
    assert len(group_members["Nhóm 2"]) == 4
    assert len(names) == 17  # Trần Mạnh Cường is a non-scoring dispatcher.


def test_canonical_member_name_accepts_unique_vietnamese_tone_typo():
    rules = [
        ServiceRule(
            3, "Fiber", Decimal("17"), "",
            (Member("Đặng Văn Minh", "Nhóm 1", "Chính"),),
        )
    ]
    assert canonical_member_name(rules, "Đăng Văn Minh") == "Đặng Văn Minh"


def test_canonical_member_name_resolves_ly_bich_hang_sheet_alias():
    rules = [
        ServiceRule(
            3, "Fiber", Decimal("17"), "",
            (Member("Lý Thị Bích Hằng", "Nhóm 1", "Chính"),),
        )
    ]

    assert canonical_member_name(rules, "Lý Bích Hằng") == "Lý Thị Bích Hằng"
    assert member_group(rules, "Lý Bích Hằng") == "Nhóm 1"


def test_loads_and_prioritizes_named_project_rules(tmp_path):
    path = tmp_path / "projects.toml"
    path.write_text(
        '''
[[projects]]
name = "Dự án BCA"
priority = 10
match_fields = ["customer_name", "subscriber_name"]
contains = "Cục Viễn Thông & Cơ Yếu Bca"
route_field = "labor_address"
[[projects.routes]]
assignee = "Đào Anh Vũ"
locations = ["Hà Nội"]

[[projects]]
name = "Đài THVN"
priority = 100
match_fields = ["subscriber_name"]
contains = "Đài THVN"
fixed_assignee = "Nguyễn Hoàng Dương"
''',
        encoding="utf-8",
    )
    rules = load_project_rules(path)
    ticket = Ticket(
        "GD1", "TB1", "Fiber", subscriber_name="Dịch vụ của ĐÀI THVN tại Hà Nội"
    )
    match = match_project_rule(rules, ticket)
    assert match is not None
    assert match.project_name == "Đài THVN"
    assert match.assignee == "Nguyễn Hoàng Dương"


def test_hanoi_radio_television_customer_is_assigned_to_nguyen_hoang_duong():
    rules = load_project_rules(Path(__file__).parents[1] / "project_rules.toml")
    ticket = Ticket(
        "GD-HNTH", "TB-HNTH", "Fiber",
        customer_name="ĐÀI PHÁT THANH - TRUYỀN HÌNH HÀ NỘI",
    )

    match = match_project_rule(rules, ticket)

    assert match is not None
    assert match.project_name == "Đài Phát Thanh - Truyền Hình Hà Nội"
    assert match.assignee == "Nguyễn Hoàng Dương"


@pytest.mark.parametrize(
    ("address", "assignee"),
    [
        ("Thành phố Hải Phòng", "Vũ Thế Ninh"),
        ("Tỉnh Điện Biên", "Vũ Thế Ninh"),
        ("TP Hà Nội", "Đào Anh Vũ"),
        ("Tỉnh Lai Châu", "Đào Anh Vũ"),
        ("Tỉnh Sơn La", "Đoàn Hải Hà"),
        ("Tỉnh Vĩnh Phúc", "Đoàn Hải Hà"),
        ("Tỉnh Hà Giang", "Nguyễn Duy Thành"),
        ("Tỉnh Thanh Hoá", "Nguyễn Duy Thành"),
        ("Tỉnh Nghệ An", "Nguyễn Duy Thành"),
        ("Tỉnh Hà Tĩnh", "Vũ Thế Ninh"),
    ],
)
def test_bca_routes_by_labor_address(address, assignee):
    rule = ProjectRule(
        name="Dự án BCA",
        contains="Cục Viễn Thông & Cơ Yếu Bca",
        match_fields=("customer_name", "subscriber_name"),
        priority=90,
        route_field="labor_address",
        routes=(
            ProjectRoute("Vũ Thế Ninh", ("Hải Phòng", "Điện Biên", "Hà Tĩnh")),
            ProjectRoute("Đào Anh Vũ", ("Hà Nội", "Lai Châu")),
            ProjectRoute("Đoàn Hải Hà", ("Sơn La", "Vĩnh Phúc")),
            ProjectRoute("Nguyễn Duy Thành", ("Hà Giang", "Thanh Hóa", "Nghệ An")),
        ),
    )
    ticket = Ticket(
        "GD1",
        "TB1",
        "Fiber",
        customer_name="CỤC VIỄN THÔNG & CƠ YẾU BCA",
        labor_address=address,
    )
    match = match_project_rule([rule], ticket)
    assert match is not None
    assert match.project_name == "Dự án BCA"
    assert match.assignee == assignee


@pytest.mark.parametrize(
    ("address", "assignee"),
    [
        ("Thành phố Hải Phòng", "Vũ Thế Ninh"),
        ("Tỉnh Hà Nội", "Đào Anh Vũ"),
        ("Tỉnh Phú Thọ", "Đoàn Hải Hà"),
        ("Tỉnh Nghệ An", "Nguyễn Duy Thành"),
    ],
)
def test_bhxh_project_uses_bca_routing(address, assignee):
    project_file = Path(__file__).parents[1] / "project_rules.toml"
    rules = load_project_rules(project_file)
    ticket = Ticket(
        "GD-BHXH", "TB-BHXH", "Fiber",
        customer_name="Ban Quản Lý Đầu Tư Và Xây Dựng Ngành Bảo Hiểm Xã Hội",
        labor_address=address,
    )

    match = match_project_rule(rules, ticket)

    assert match is not None
    assert match.project_name == "Dự án BHXH"
    assert match.assignee == assignee


def test_bca_unknown_address_is_rejected_instead_of_using_service_rule():
    rule = ProjectRule(
        name="Dự án BCA",
        contains="Cục Viễn Thông & Cơ Yếu Bca",
        match_fields=("customer_name",),
        priority=90,
        route_field="labor_address",
        routes=(ProjectRoute("Đào Anh Vũ", ("Hà Nội",)),),
    )
    ticket = Ticket(
        "GD1", "TB1", "Fiber", customer_name="Cục Viễn Thông & Cơ Yếu Bca"
    )
    with pytest.raises(RuleError, match="Dự án BCA"):
        match_project_rule([rule], ticket)


def test_mobifone_no_longer_uses_project_routing():
    project_file = Path(__file__).parents[1] / "project_rules.toml"
    rules = load_project_rules(project_file)
    ticket = Ticket(
        "GD-MOBI", "TB-MOBI", "Kênh thuê riêng",
        customer_name=(
            "TRUNG TÂM QUẢN LÝ, ĐIỀU HÀNH MẠNG (NOC) - CHI NHÁNH "
            "TỔNG CÔNG TY VIỄN THÔNG MOBIFONE"
        ),
        labor_address="Tòa nhà Mobifone, Thành phố Hà Nội",
    )
    assert match_project_rule(rules, ticket) is None


@pytest.mark.parametrize(
    ("address", "assignee"),
    [
        ("Điểm lắp đặt tại tỉnh Thái Bình", "Vũ Thế Ninh"),
        ("Địa chỉ LĐ, Thành phố Bắc Ninh", "Đào Anh Vũ"),
        ("Chi nhánh tại tỉnh Yên Bái", "Đoàn Hải Hà"),
        ("Trạm thiết bị tỉnh Lạng Sơn", "Nguyễn Duy Thành"),
        ("Điểm đặt tại Nghệ An", "Nguyễn Duy Thành"),
    ],
)
def test_bdtw_project_uses_bca_routing(address, assignee):
    project_file = Path(__file__).parents[1] / "project_rules.toml"
    rules = load_project_rules(project_file)
    ticket = Ticket(
        "GD-BDTW", "TB-BDTW", "Fiber",
        customer_name="Trung tâm dịch vụ thuộc CỤC BĐTW",
        labor_address=address,
    )
    match = match_project_rule(rules, ticket)
    assert match is not None
    assert match.project_name == "Dự án Cục BĐTW"
    assert match.assignee == assignee


@pytest.mark.parametrize(
    ("address", "assignee"),
    [
        ("Điểm lắp đặt tại tỉnh Hải Dương", "Vũ Thế Ninh"),
        ("Địa chỉ LĐ, Thành phố Hà Nội", "Đào Anh Vũ"),
        ("Chi nhánh tại tỉnh Phú Thọ", "Đoàn Hải Hà"),
        ("Trạm thiết bị tỉnh Quảng Ninh", "Nguyễn Duy Thành"),
        ("Điểm đặt tại Hà Tĩnh", "Vũ Thế Ninh"),
    ],
)
def test_btc_project_uses_bca_routing(address, assignee):
    project_file = Path(__file__).parents[1] / "project_rules.toml"
    rules = load_project_rules(project_file)
    ticket = Ticket(
        "GD-BTC", "TB-BTC", "Fiber",
        customer_name="Đơn vị trực thuộc TỔNG CỤC DỰ TRỮ NHÀ NƯỚC",
        labor_address=address,
    )
    match = match_project_rule(rules, ticket)
    assert match is not None
    assert match.project_name == "Dự án BTC"
    assert match.assignee == assignee


def test_btc_project_accepts_tong_cuc_thue_customer_alias():
    project_file = Path(__file__).parents[1] / "project_rules.toml"
    rules = load_project_rules(project_file)
    ticket = Ticket(
        "GD-THUE", "TB-THUE", "Fiber",
        customer_name="Đơn vị trực thuộc TỔNG CỤC THUẾ",
        labor_address="Chi cục Thuế Kỳ Sơn",
        connection_address="Cục Thuế TP. Hải Phòng, Tỉnh Hải Phòng",
        labor_province="Thành phố Hà Nội",
    )
    match = match_project_rule(rules, ticket)
    assert match is not None
    assert match.project_name == "Dự án BTC"
    assert match.assignee == "Đào Anh Vũ"


def test_btc_project_accepts_customs_it_statistics_customer_alias():
    project_file = Path(__file__).parents[1] / "project_rules.toml"
    rules = load_project_rules(project_file)
    ticket = Ticket(
        "GD-CUSTOMS", "TB-CUSTOMS", "Megawan",
        customer_name="Cục Công Nghệ Thông Tin & Thống Kê Hải Quan",
        labor_address="Điểm lắp đặt tại tỉnh Thanh Hóa",
    )
    match = match_project_rule(rules, ticket)
    assert match is not None
    assert match.project_name == "Dự án BTC"
    assert match.assignee == "Nguyễn Duy Thành"


def test_vietlott_project_routes_tuyen_quang_by_subscriber_prefix(tmp_path):
    path = tmp_path / "projects.toml"
    path.write_text(
        '''
[[projects]]
name = "Dự án Vietlott"
priority = 90
match_fields = ["customer_name"]
contains = "Công Ty Cổ Phần Đầu Tư Kỹ Thuật Berjaya Gia Thịnh"
route_field = "labor_address"

[[projects.routes]]
assignee = "Đào Anh Vũ"
locations = ["Tuyên Quang"]
subscriber_prefixes = ["bgt5"]

[[projects.routes]]
assignee = "Nguyễn Duy Thành"
locations = ["Tuyên Quang"]
subscriber_prefixes = ["bgt2"]
''',
        encoding="utf-8",
    )
    rules = load_project_rules(path)
    base = dict(
        transaction_id="VNP-LD/00078532",
        service="Fiber",
        customer_name="Công Ty Cổ Phần Đầu Tư Kỹ Thuật Berjaya Gia Thịnh",
        labor_address="Thành phố Tuyên Quang",
    )
    assert match_project_rule(
        rules, Ticket(subscriber_id="bgt57174", **base)
    ).assignee == "Đào Anh Vũ"
    assert match_project_rule(
        rules, Ticket(subscriber_id="bgt27174", **base)
    ).assignee == "Nguyễn Duy Thành"


@pytest.mark.parametrize(
    ("address", "assignee"),
    [
        ("Điểm lắp đặt tại Hà Nội", "Đào Anh Vũ"),
        ("Điểm lắp đặt tại Lạng Sơn", "Đoàn Hải Hà"),
        ("Điểm lắp đặt tại Thanh Hóa", "Nguyễn Duy Thành"),
        ("Điểm lắp đặt tại Hà Giang", "Vũ Thế Ninh"),
        ("Điểm lắp đặt tại Hà Tĩnh", "Vũ Thế Ninh"),
    ],
)
def test_vietlott_reassigned_routes(address, assignee):
    project_file = Path(__file__).parents[1] / "project_rules.toml"
    rules = load_project_rules(project_file)
    ticket = Ticket(
        "GD-VIETLOTT", "BGT0001", "Fiber",
        customer_name="Công Ty Cổ Phần Đầu Tư Kỹ Thuật Berjaya Gia Thịnh",
        labor_address=address,
    )

    match = match_project_rule(rules, ticket)

    assert match is not None
    assert match.project_name == "Dự án Vietlott"
    assert match.assignee == assignee


def test_le_duc_vinh_is_not_assigned_to_any_project():
    rules = load_project_rules(Path(__file__).parents[1] / "project_rules.toml")

    assert all(
        route.assignee != "Lê Đức Vinh"
        for project in rules
        for route in project.routes
    )
    assert all(project.fixed_assignee != "Lê Đức Vinh" for project in rules)


def test_project_falls_back_to_labor_province():
    project_file = Path(__file__).parents[1] / "project_rules.toml"
    rules = load_project_rules(project_file)
    ticket = Ticket(
        "VNP-TD/00097957", "MW000020934", "Megawan",
        customer_name="Tổng Cục Dự Trữ Nhà Nước",
        labor_address="Chi cục Thuế Cát Hải",
        connection_address="Địa chỉ kết nối tại Thành phố Hà Nội",
        labor_province="Tp Hải Phòng",
    )
    match = match_project_rule(rules, ticket)
    assert match is not None
    assert match.project_name == "Dự án BTC"
    assert match.assignee == "Vũ Thế Ninh"


def test_kho_bac_nha_nuoc_uses_btc_project_routes():
    project_file = Path(__file__).parents[1] / "project_rules.toml"
    rules = load_project_rules(project_file)
    ticket = Ticket(
        "VNP-TD/00098169", "MW000020934", "Megawan",
        customer_name="Kho Bạc Nhà Nước tỉnh Hải Phòng",
        labor_address="Trụ sở Kho Bạc Nhà Nước, Thành phố Hải Phòng",
    )

    match = match_project_rule(rules, ticket)

    assert match is not None
    assert match.project_name == "Dự án BTC"
    assert match.assignee == "Vũ Thế Ninh"


def test_hnth_customer_rule_requires_matching_note():
    rules = load_project_rules(Path(__file__).parents[1] / "project_rules.toml")
    ticket = Ticket(
        "VNP-LD/00078688", "mwk000zyh", "Megawan",
        customer_name="Cục Quản Trị Ngân Hàng Nhà Nước Việt Nam",
        notes="Lắp đặt Kênh phục vụ HNTH theo yêu cầu",
    )

    match = match_project_rule(rules, ticket)

    assert match is not None
    assert match.project_name == "Cục Quản Trị NHNN - Kênh phục vụ HNTH"
    assert match.assignee == "Nguyễn Hoàng Dương"
    assert match_project_rule(
        rules,
        Ticket(
            "VNP-LD/00078688", "mwk000zyh", "Megawan",
            customer_name="Cục Quản Trị Ngân Hàng Nhà Nước Việt Nam",
            notes="Kênh thông thường",
        ),
    ) is None


def test_project_understands_tp_bg_as_bac_giang():
    project_file = Path(__file__).parents[1] / "project_rules.toml"
    rules = load_project_rules(project_file)
    ticket = Ticket(
        "VNP-TD/00097966", "MW000021095_1", "Megawan quang GE",
        customer_name="Tổng Cục Dự Trữ Nhà Nước",
        labor_address="303 Lê Lợi TP BG (Nội Tỉnh)",
    )
    match = match_project_rule(rules, ticket)
    assert match is not None
    assert match.project_name == "Dự án BTC"
    assert match.assignee == "Đào Anh Vũ"


def test_tong_cuc_thue_falls_back_to_labor_province():
    project_file = Path(__file__).parents[1] / "project_rules.toml"
    rules = load_project_rules(project_file)
    ticket = Ticket(
        "VNP-TD/00097982", "MW000020500", "Megawan",
        customer_name="Tổng Cục Thuế",
        labor_address="Chi cục Thuế Kỳ Sơn",
        connection_address=(
            "Cục Thuế TP. Hải Phòng, Phường Vĩnh Niệm, Tỉnh Hải Phòng"
        ),
        labor_province="Tỉnh Nghệ An",
    )

    match = match_project_rule(rules, ticket)

    assert match is not None
    assert match.project_name == "Dự án BTC"
    assert match.assignee == "Nguyễn Duy Thành"


def test_project_prefers_labor_address_over_labor_province():
    project_file = Path(__file__).parents[1] / "project_rules.toml"
    rules = load_project_rules(project_file)
    ticket = Ticket(
        "GD-THUE", "TB-THUE", "Megawan",
        customer_name="Tổng Cục Thuế",
        labor_address="Điểm lắp đặt tại Thành phố Hải Phòng",
        connection_address="Cục Thuế TP. Hải Phòng, Tỉnh Hải Phòng",
        labor_province="Tỉnh Nghệ An",
    )

    match = match_project_rule(rules, ticket)

    assert match is not None
    assert match.assignee == "Vũ Thế Ninh"


def test_project_does_not_use_connection_address_as_fallback():
    project_file = Path(__file__).parents[1] / "project_rules.toml"
    rules = load_project_rules(project_file)
    ticket = Ticket(
        "GD-THUE", "TB-THUE", "Megawan",
        customer_name="Tổng Cục Thuế",
        labor_address="Chi cục Thuế không ghi tỉnh",
        connection_address="Cục Thuế TP. Hải Phòng, Tỉnh Hải Phòng",
    )

    with pytest.raises(RuleError, match="labor_province"):
        match_project_rule(rules, ticket)
